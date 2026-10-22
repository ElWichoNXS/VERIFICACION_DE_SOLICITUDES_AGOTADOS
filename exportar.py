"""Genera el Excel de resultados con formato."""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from motor import Resultado

MORADO = "5B2C83"
ROJO = "C00000"
thin = Side(style="thin", color="BFBFBF")
BORDE = Border(left=thin, right=thin, top=thin, bottom=thin)
F_ROJO = PatternFill("solid", fgColor="F8CBAD")
F_NAR = PatternFill("solid", fgColor="FCE4D6")
F_VERDE = PatternFill("solid", fgColor="E2EFDA")

ANCHOS = {
    "Descripción": 42, "Motivo": 48, "Destinos": 22, "2. Alternativos": 24, "4. Extremos": 24,
    "Tipo de Origen": 30, "Validación": 22, "Supervisor Origen": 16, "Supervisor": 16, "Nombre Local": 22,
    "Filas Solicitud": 12, "Repetido": 14, "Rotación Origen": 14,
}


def _hoja(wb: Workbook, nombre: str, df: pd.DataFrame, titulo: str, nota: str, color: str = MORADO):
    ws = wb.create_sheet(nombre)
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=14, color=color)
    ws["A2"] = nota
    ws["A2"].font = Font(italic=True, size=9, color="595959")
    fila_h = 4
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(fila_h, j, col)
        c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
        ws.column_dimensions[get_column_letter(j)].width = ANCHOS.get(col, max(10, min(16, len(str(col)) + 2)))
    ws.row_dimensions[fila_h].height = 40
    for i, fila in enumerate(df.itertuples(index=False), fila_h + 1):
        for j, v in enumerate(fila, 1):
            if isinstance(v, float):
                if pd.isna(v):
                    v = None
                elif v.is_integer():
                    v = int(v)
            c = ws.cell(i, j, v)
            c.border = BORDE
            if df.columns[j - 1] not in ("Descripción", "Motivo", "Destinos", "2. Alternativos", "4. Extremos",
                                         "Tipo de Origen", "Nombre Local"):
                c.alignment = Alignment(horizontal="center")
    ultima = fila_h + max(len(df), 1)
    if len(df.columns):
        ws.auto_filter.ref = f"A{fila_h}:{get_column_letter(len(df.columns))}{ultima}"
    ws.freeze_panes = ws.cell(fila_h + 1, 1)
    return ws, fila_h, ultima


def _col(df: pd.DataFrame, nombre: str) -> str | None:
    return get_column_letter(list(df.columns).index(nombre) + 1) if nombre in df.columns else None


def _resaltar_mayor_cero(ws, df, nombre, ini, fin, fill=F_NAR):
    c = _col(df, nombre)
    if c:
        ws.conditional_formatting.add(f"{c}{ini}:{c}{fin}", CellIsRule(operator="greaterThan", formula=["0"],
                                      fill=fill, font=Font(bold=True, color=ROJO)))


def generar_excel(res: Resultado) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    r = res.resumen

    # 1) TRASPASOS A GRABAR
    tg = res.traspasos
    ws, h, u = _hoja(wb, "TRASPASOS A GRABAR", tg, "TRASPASOS A GRABAR",
                     "Lista consolidada. Columnas A:D en el orden para grabar en el sistema (origen, destino, estadístico, cantidad).")
    c = _col(tg, "¿Queda Bajo Exhibición?")
    if c:
        ws.conditional_formatting.add(f"{c}{h+1}:{c}{u}", CellIsRule(operator="equal", formula=['"SÍ"'],
                                      fill=F_ROJO, font=Font(bold=True, color=ROJO)))
    _resaltar_mayor_cero(ws, tg, "Und. Bajo Exhibición Causadas", h + 1, u)

    # 2) REPORTE BAJO EXHIBICIÓN
    ws, h, u = _hoja(wb, "REPORTE BAJO EXHIBICIÓN", res.reporte, "REPORTE: LOCALES QUE QUEDAN BAJO EXHIBICIÓN",
                     "Un registro por local e ítem que, después de todos los traspasos, queda por debajo de su exhibición.", ROJO)
    _resaltar_mayor_cero(ws, res.reporte, "Und. Bajo Exhibición Causadas", h + 1, u)

    # 3) DETALLE POR SOLICITUD
    det = res.detalle
    ws, h, u = _hoja(wb, "DETALLE POR SOLICITUD", det, "DETALLE POR SOLICITUD",
                     "Cómo se cubrió cada solicitud: 1) origen sin bajar exhibición, 2) alternativos con sobrante, "
                     "3) origen bajo exhibición, 4) otros locales bajo exhibición. Entre paréntesis, las unidades.")
    c = _col(det, "Validación")
    if c:
        for val, fill, color in [("OK", F_VERDE, "375623"), ("DEJA BAJO EXHIBICIÓN", F_NAR, ROJO),
                                 ("STOCK INSUFICIENTE", F_ROJO, ROJO), ("SIN STOCK EN ORIGEN", F_ROJO, ROJO),
                                 ("SIN DATO EN BASE", PatternFill("solid", fgColor="EDEDED"), "595959")]:
            ws.conditional_formatting.add(f"{c}{h+1}:{c}{u}", CellIsRule(operator="equal", formula=[f'"{val}"'],
                                          fill=fill, font=Font(bold=True, color=color)))
    c = _col(det, "Repetido")
    if c:
        ws.conditional_formatting.add(f"{c}{h+1}:{c}{u}", CellIsRule(operator="equal", formula=['"REPETIDO"'],
                                      fill=F_ROJO, font=Font(bold=True, color=ROJO)))
    for n in ["3. Origen bajo exhibición", "Und. Extremos"]:
        _resaltar_mayor_cero(ws, det, n, h + 1, u)
    _resaltar_mayor_cero(ws, det, "Pendiente sin Cubrir", h + 1, u, F_ROJO)

    # 4) PENDIENTES (si hay)
    if len(res.pendientes):
        _hoja(wb, "PENDIENTES", res.pendientes, "UNIDADES SIN CUBRIR",
              "No hay stock en ningún local de la base para completar estas solicitudes.", ROJO)

    # 5) RESUMEN (primera hoja)
    ws = wb.create_sheet("RESUMEN", 0)
    ws["A1"] = "RESUMEN DE VERIFICACIÓN DE TRASPASOS"
    ws["A1"].font = Font(bold=True, size=14, color=MORADO)
    ws["A2"] = f"Generado: {datetime.now():%d/%m/%Y %H:%M}"
    ws["A2"].font = Font(italic=True, size=9, color="595959")
    filas = [
        ("SOLICITUDES", None, True),
        ("Solicitudes", r["solicitudes"], False),
        ("Unidades solicitadas", r["und_solicitadas"], False),
        ("Filas repetidas (mismo origen, destino e ítem)", r["repetidos"], False),
        ("", None, False),
        ("¿DE DÓNDE SALEN LAS UNIDADES?", None, True),
        ("1. Origen solicitado, sin bajar de exhibición", r["paso1"], False),
        ("2. Locales alternativos con sobrante", r["paso2"], False),
        ("3. Origen solicitado, quedando bajo exhibición", r["paso3"], False),
        ("4. Otros locales, quedando bajo exhibición (extremos)", r["paso4"], False),
        ("Pendiente sin cubrir", r["pendiente"], False),
        ("", None, False),
        ("RESULTADO", None, True),
        ("Traspasos a grabar", r["traspasos"], False),
        ("Locales/ítems que quedan bajo exhibición", r["locales_bajo_exh"], False),
        ("", None, False),
        ("VALIDACIÓN INICIAL DE LAS SOLICITUDES", None, True),
    ] + [(k, v, False) for k, v in r["validacion"].items()]
    for i, (k, v, t) in enumerate(filas, 4):
        a = ws.cell(i, 1, k)
        if t:
            a.font = Font(bold=True, color=MORADO)
        if v is not None:
            ws.cell(i, 2, int(v) if float(v).is_integer() else v).font = Font(bold=True)
    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
